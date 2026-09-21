"""Minta HR dokumentumok a RAG teszthez."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "sample_docs"

HR_DOC = """A vállalat neve: Example Kft.

A munkarend:
Hétfő–péntek 8:00–16:00.

Az éves szabadság:
25 munkanap.

3. fejezet
A próbaidő 3 hónap.

A szerződés időtartama: határozatlan idejű.
Az ügyfél kötelezettségei: a szabályzat betartása.
A szolgáltatás igénybevételének feltétele: érvényes munkaviszony.
"""


def make_pdf(text: str) -> bytes:
    commands = ["BT", "/F1 12 Tf"]
    y = 720
    for line in text.splitlines() or [text]:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"72 {y} Td ({safe}) Tj")
        commands.append("0 -18 Td")
        y -= 18
    commands.append("ET")
    stream = "\n".join(commands).encode("latin-1", errors="replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode()
        out += obj
        out += b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Munkaszabályzat", level=1)
    for paragraph in text.split("\n\n"):
        document.add_paragraph(paragraph)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Név"
    table.rows[0].cells[1].text = "Beosztás"
    table.rows[1].cells[0].text = "Kiss Péter"
    table.rows[1].cells[1].text = "Fejlesztő"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dolgozók"
    sheet.append(["Név", "Beosztás", "Fizetés"])
    sheet.append(["Kiss Péter", "Fejlesztő", 650000])
    sheet.append(["Nagy Anna", "Projektvezető", 900000])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "munkaszabalyzat.txt").write_text(HR_DOC, encoding="utf-8")
    (OUT / "munkaszabalyzat.pdf").write_bytes(make_pdf(HR_DOC))
    (OUT / "munkaszabalyzat.docx").write_bytes(make_docx(HR_DOC))
    (OUT / "dolgozok.xlsx").write_bytes(make_xlsx())
    print(f"Sample documents written to {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
