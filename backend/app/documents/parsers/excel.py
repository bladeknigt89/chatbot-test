from pathlib import Path

from app.documents.parsers.base import ExtractedBlock, ExtractedDocument


def parse_excel(path: Path) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return _parse_xls(path)
    return _parse_xlsx(path)


def _parse_xlsx(path: Path) -> ExtractedDocument:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), data_only=True, read_only=True)
    blocks: list[ExtractedBlock] = []
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            sheet_text = _sheet_rows_to_text(_iter_xlsx_rows(sheet), sheet.title)
            if not sheet_text:
                continue
            blocks.append(ExtractedBlock(text=sheet_text, sheet_name=sheet.title))
            parts.append(sheet_text)
    finally:
        workbook.close()
    return _finish(parts, blocks)


def _iter_xlsx_rows(sheet):
    for row in sheet.iter_rows(values_only=True):
        yield row


def _parse_xls(path: Path) -> ExtractedDocument:
    import xlrd

    book = xlrd.open_workbook(str(path))
    blocks: list[ExtractedBlock] = []
    parts: list[str] = []
    for sheet in book.sheets():
        rows = (sheet.row_values(idx) for idx in range(sheet.nrows))
        sheet_text = _sheet_rows_to_text(rows, sheet.name)
        if not sheet_text:
            continue
        blocks.append(ExtractedBlock(text=sheet_text, sheet_name=sheet.name))
        parts.append(sheet_text)
    return _finish(parts, blocks)


def _sheet_rows_to_text(rows, title: str) -> str:
    rendered: list[str] = [f"Munkalap: {title}"]
    for row in rows:
        values = [_stringify(value) for value in row]
        if not any(values):
            continue
        rendered.append(" | ".join(values))
    if len(rendered) == 1:
        return ""
    return "\n".join(rendered)


def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _finish(parts: list[str], blocks: list[ExtractedBlock]) -> ExtractedDocument:
    joined = "\n\n".join(parts).strip()
    if not joined:
        raise ValueError("Az Excel fájl nem tartalmaz kinyerhető adatot.")
    return ExtractedDocument(text=joined, blocks=blocks)
