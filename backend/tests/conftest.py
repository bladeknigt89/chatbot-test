from __future__ import annotations

import hashlib
import os
from collections.abc import Iterator
from pathlib import Path

import numpy as np
import pytest
from fastapi.testclient import TestClient

os.environ.setdefault("APP_ENV", "test")
os.environ.setdefault("LLM_SKIP_CHECK", "true")
os.environ.setdefault("ADMIN_USERNAME", "ai")
os.environ.setdefault("ADMIN_PASSWORD", "No_comment_123")
os.environ.setdefault("SECRET_KEY", "test-secret-key-not-for-production")


class HashEmbedding:
    dim = 64

    def embed_query(self, text: str) -> list[float]:
        return self._vec(text)

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vec(text) for text in texts]

    def _vec(self, text: str) -> list[float]:
        vec = np.zeros(self.dim, dtype=np.float32)
        for token in text.lower().replace("—", " ").replace("–", " ").split():
            token = "".join(ch for ch in token if ch.isalnum())
            if not token:
                continue
            digest = hashlib.md5(token.encode("utf-8")).digest()
            idx = int.from_bytes(digest[:2], "little") % self.dim
            vec[idx] += 1.0
        norm = float(np.linalg.norm(vec) or 1.0)
        return (vec / norm).tolist()


class ExtractiveLLM:
    def generate(self, messages: list[dict[str, str]], temperature: float) -> str:
        from app.rag.prompts import no_info_reply

        user = messages[-1]["content"]
        question = user.split("User question:")[-1]
        question = question.split("Answer only")[0].strip()
        if "No document excerpts were retrieved" in user:
            return no_info_reply(question)
        context = user.split("Document excerpts:")[-1].split("User question:")[0]
        lowered_q = question.lower()
        lowered_c = context.lower()
        if any(word in lowered_q for word in ["árbevétel", "revenue", "profit", "nyereség"]):
            if not any(word in lowered_c for word in ["árbevétel", "revenue", "profit", "nyereség"]):
                return no_info_reply(question)
        if "3. fejezet" in lowered_q or "fejezetében" in lowered_q:
            if "próbaidő" in lowered_c or "probaido" in lowered_c:
                return "A 3. fejezet szerint a próbaidő 3 hónap."
        if "foglald össze" in lowered_q or "foglalja össze" in lowered_q:
            return "A dokumentum az Example Kft. munkarendjét, szabadságát és szerződéses feltételeit írja le."
        if "feltételekkel" in lowered_q or "igénybe" in lowered_q:
            if "munkaviszony" in lowered_c:
                return "A szolgáltatás igénybevételének feltétele: érvényes munkaviszony."
        if "időtartama" in lowered_q:
            if "határozatlan" in lowered_c:
                return "A szerződés időtartama határozatlan idejű."
        if "kötelezettségei" in lowered_q or "ugyfelnek" in lowered_q:
            if "szabályzat" in lowered_c or "szabalyzat" in lowered_c:
                return "Az ügyfél kötelezettségei: a szabályzat betartása."
        sentences = [
            part.strip()
            for part in context.replace("\n", ". ").split(". ")
            if part.strip() and not part.strip().startswith("[dokumentum=")
        ]
        q_tokens = {tok for tok in lowered_q.split() if len(tok) > 3}
        ranked = []
        for sentence in sentences:
            score = sum(1 for tok in q_tokens if tok in sentence.lower())
            ranked.append((score, sentence))
        ranked.sort(reverse=True)
        if ranked and ranked[0][0] > 0:
            return ranked[0][1].strip() + "."
        if "Example Kft" in context:
            return context.strip()[:500]
        return no_info_reply(question)

    def stream(self, messages: list[dict[str, str]], temperature: float) -> Iterator[str]:
        yield self.generate(messages, temperature)


def make_pdf(text: str) -> bytes:
    commands = ["BT", "/F1 12 Tf"]
    y = 720
    for line in text.splitlines() or [text]:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"72 {y} Td ({safe}) Tj")
        commands.append("0 -18 Td")
        y -= 18
    commands.append("ET")
    stream = "\n".join(commands).encode("latin-1", errors="replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode()
        out += obj
        out += b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


def make_docx(paragraphs: list[str], tables: list[list[list[str]]] | None = None) -> bytes:
    from io import BytesIO

    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for table_data in tables or []:
        table = document.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, value in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, list[list[str]]]) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = name
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


HR_DOC = """A vállalat neve: Example Kft.

A munkarend:
Hétfő–péntek 8:00–16:00.

Az éves szabadság:
25 munkanap.

3. fejezet
A próbaidő 3 hónap.

A szerződés időtartama: határozatlan idejű.
Az ügyfél kötelezettségei: a szabályzat betartása.
A szolgáltatás igénybevételének feltétele: érvényes munkaviszony.
"""


@pytest.fixture()
def app_env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    db_path = tmp_path / "app.db"
    vector_path = tmp_path / "vector.db"
    storage = tmp_path / "uploads"
    logs = tmp_path / "logs"
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.setenv("LLM_SKIP_CHECK", "true")
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_path.as_posix()}")
    monkeypatch.setenv("VECTOR_DB_PATH", str(vector_path))
    monkeypatch.setenv("STORAGE_PATH", str(storage))
    monkeypatch.setenv("LOG_PATH", str(logs))
    monkeypatch.setenv("CHAT_HISTORY_ENABLED", "true")
    monkeypatch.setenv("MAX_FILE_SIZE", "1048576")
    from app.config import get_settings
    from app.embeddings import factory as embedding_factory
    from app.llm import factory as llm_factory
    from app.vectorstore.store import reset_vector_store

    get_settings.cache_clear()
    embedding_factory.reset_embedding_provider()
    llm_factory.reset_llm_provider()
    reset_vector_store()

    fake_embed = HashEmbedding()
    fake_llm = ExtractiveLLM()
    monkeypatch.setattr(embedding_factory, "get_embedding_provider", lambda: fake_embed)
    monkeypatch.setattr(llm_factory, "get_llm_provider", lambda: fake_llm)

    from app.api import chat as chat_api
    from app.main import create_app
    from app.workers import processor

    monkeypatch.setattr(processor, "get_embedding_provider", lambda: fake_embed)
    monkeypatch.setattr(chat_api, "get_embedding_provider", lambda: fake_embed)
    monkeypatch.setattr(chat_api, "get_llm_provider", lambda: fake_llm)

    application = create_app()
    return application, fake_embed, fake_llm


@pytest.fixture()
def client(app_env) -> TestClient:
    application, _, _ = app_env
    with TestClient(application) as test_client:
        yield test_client


@pytest.fixture()
def admin_client(client: TestClient) -> TestClient:
    response = client.post("/api/auth/login", json={"username": "ai", "password": "No_comment_123"})
    assert response.status_code == 200, response.text
    return client


def drain_jobs() -> None:
    from app.database import get_session_factory
    from app.workers.processor import process_job
    from app.workers.queue import claim_next_job

    db = get_session_factory()()
    try:
        while True:
            job = claim_next_job(db)
            if not job:
                break
            process_job(db, job)
    finally:
        db.close()
